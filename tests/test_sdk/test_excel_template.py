from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from sdk.excel_template import fill_excel_template, scan_excel_placeholders


def _create_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "订单号"
    sheet["B1"] = "{{order_no}}"
    sheet["B1"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    sheet["A2"] = "数量"
    sheet["B2"] = "{{quantity}} 件"
    sheet["C2"] = "{{missing_value}}"
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "签收人：{{signer_name}}"
    workbook.save(path)


def test_scan_excel_placeholders_finds_field_keys_and_locations(tmp_path):
    template_path = tmp_path / "template.xlsx"
    _create_template(template_path)

    placeholders = scan_excel_placeholders(template_path)

    assert [
        (item.sheet_name, item.coordinate, item.field_key, item.raw_value)
        for item in placeholders
    ] == [
        ("Report", "B1", "order_no", "{{order_no}}"),
        ("Report", "B2", "quantity", "{{quantity}} 件"),
        ("Report", "C2", "missing_value", "{{missing_value}}"),
        ("Report", "A4", "signer_name", "签收人：{{signer_name}}"),
    ]


def test_fill_excel_template_writes_extracted_values_and_preserves_style(tmp_path):
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "filled.xlsx"
    _create_template(template_path)

    result = fill_excel_template(
        template_path,
        output_path,
        {
            "order_no": "NOZS0311046",
            "quantity": 0,
            "signer_name": "张三",
        },
    )

    assert result.replaced_count == 3
    assert result.missing_fields == ["missing_value"]

    workbook = load_workbook(output_path)
    sheet = workbook["Report"]
    assert sheet["B1"].value == "NOZS0311046"
    assert sheet["B1"].fill.fgColor.rgb == "00FFF2CC"
    assert sheet["B2"].value == "0 件"
    assert sheet["C2"].value is None
    assert sheet["A4"].value == "签收人：张三"
    assert "A4:B4" in {str(item) for item in sheet.merged_cells.ranges}

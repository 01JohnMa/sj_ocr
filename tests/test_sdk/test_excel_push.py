from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch
import importlib.util

import pytest
from openpyxl import Workbook, load_workbook


DOCUMENT_ID = "33333333-3333-4333-8333-333333333333"
TEMPLATE_ID = "44444444-4444-4444-8444-444444444444"


def _load_document_helpers_module():
    module_path = Path(__file__).resolve().parents[2] / "api" / "routes" / "documents" / "helpers.py"
    spec = importlib.util.spec_from_file_location("document_helpers_under_test", module_path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


helpers = _load_document_helpers_module()


@pytest.mark.asyncio
async def test_push_to_feishu_attaches_filled_excel_template(tmp_path, monkeypatch):
    """配置固定 Excel 模板时，推送附件应包含填值后的 xlsx。"""
    template_path = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "订单号"
    sheet["B1"] = "{{order_no}}"
    workbook.save(template_path)

    source_path = tmp_path / "source.pdf"
    source_path.write_bytes(b"%PDF-1.4")

    uploaded_paths = []

    async def fake_upload(file_path, _bitable_token):
        uploaded_paths.append(Path(file_path))
        if file_path.endswith(".xlsx"):
            uploaded_workbook = load_workbook(file_path)
            assert uploaded_workbook.active["B1"].value == "NOZS0311046"
            return "excel-token"
        return "source-token"

    fake_feishu = MagicMock()
    fake_feishu._upload_file_to_feishu = AsyncMock(side_effect=fake_upload)
    fake_feishu.push_by_template = AsyncMock(return_value=True)

    monkeypatch.setattr(helpers.settings, "UPLOAD_FOLDER", str(tmp_path / "uploads"))
    monkeypatch.setattr(helpers, "has_feishu_push_record", AsyncMock(return_value=False))
    monkeypatch.setattr(helpers, "record_feishu_push", AsyncMock())
    mock_template_service = MagicMock()
    mock_template_service.build_field_mapping.return_value = {"order_no": "订单号"}
    monkeypatch.setattr(helpers, "template_service", mock_template_service)

    with patch("services.feishu_service.feishu_service", fake_feishu):
        await helpers.push_to_feishu(
            template={
                "id": TEMPLATE_ID,
                "name": "固定模板",
                "feishu_bitable_token": "bitable-token",
                "feishu_table_id": "table-id",
                "output_mode": "both",
                "excel_template_path": str(template_path),
                "excel_template_file_name": "template.xlsx",
            },
            extraction_data={"order_no": "NOZS0311046"},
            display_name="推送文件",
            document_id=DOCUMENT_ID,
            source_file_path=str(source_path),
        )

    assert len(uploaded_paths) == 2
    assert uploaded_paths[0] == source_path
    assert uploaded_paths[1].suffix == ".xlsx"
    push_data = fake_feishu.push_by_template.await_args.args[0]
    assert push_data["attachment"] == [
        {"file_token": "source-token"},
        {"file_token": "excel-token"},
    ]


@pytest.mark.asyncio
async def test_push_to_feishu_merges_extra_crm_fields(monkeypatch):
    """CRM 手填字段应只作为额外飞书字段进入推送 payload。"""
    fake_feishu = MagicMock()
    fake_feishu.push_by_template = AsyncMock(return_value=True)

    monkeypatch.setattr(helpers, "has_feishu_push_record", AsyncMock(return_value=False))
    monkeypatch.setattr(helpers, "record_feishu_push", AsyncMock())
    mock_template_service = MagicMock()
    mock_template_service.build_field_mapping.return_value = {"sample_name": "样品名称"}
    monkeypatch.setattr(helpers, "template_service", mock_template_service)

    with patch("services.feishu_service.feishu_service", fake_feishu):
        pushed = await helpers.push_to_feishu(
            template={
                "id": TEMPLATE_ID,
                "name": "检测报告",
                "feishu_bitable_token": "bitable-token",
                "feishu_table_id": "table-id",
            },
            extraction_data={"sample_name": "LED灯"},
            display_name="推送文件",
            document_id=DOCUMENT_ID,
            extra_data={
                "alipay_account": "pay@example.com",
                "alipay_name": "张三",
            },
            extra_field_mapping={
                "alipay_account": "支付宝账号",
                "alipay_name": "支付宝姓名",
            },
        )

    assert pushed is True
    push_data, field_mapping, _, _ = fake_feishu.push_by_template.await_args.args
    assert push_data["sample_name"] == "LED灯"
    assert push_data["alipay_account"] == "pay@example.com"
    assert push_data["alipay_name"] == "张三"
    assert field_mapping["sample_name"] == "样品名称"
    assert field_mapping["alipay_account"] == "支付宝账号"
    assert field_mapping["alipay_name"] == "支付宝姓名"

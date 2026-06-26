"""Fixed Excel template placeholder scanning and filling."""

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any, Mapping

from openpyxl import load_workbook


PLACEHOLDER_PATTERN = re.compile(r"{{\s*([A-Za-z_][A-Za-z0-9_]*)\s*}}")


@dataclass(frozen=True)
class ExcelPlaceholder:
    sheet_name: str
    coordinate: str
    field_key: str
    raw_value: str


@dataclass(frozen=True)
class ExcelFillResult:
    output_path: str
    replaced_count: int
    missing_fields: list[str]
    placeholders: list[ExcelPlaceholder]


def scan_excel_placeholders(template_path: str | Path) -> list[ExcelPlaceholder]:
    workbook = load_workbook(template_path, data_only=False)
    placeholders: list[ExcelPlaceholder] = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                for match in PLACEHOLDER_PATTERN.finditer(cell.value):
                    placeholders.append(
                        ExcelPlaceholder(
                            sheet_name=sheet.title,
                            coordinate=cell.coordinate,
                            field_key=match.group(1),
                            raw_value=cell.value,
                        )
                    )

    return placeholders


def fill_excel_template(
    template_path: str | Path,
    output_path: str | Path,
    values: Mapping[str, Any],
) -> ExcelFillResult:
    workbook = load_workbook(template_path, data_only=False)
    placeholders: list[ExcelPlaceholder] = []
    missing_fields: list[str] = []
    replaced_count = 0

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue

                matches = list(PLACEHOLDER_PATTERN.finditer(cell.value))
                if not matches:
                    continue

                raw_value = cell.value
                new_value: Any = raw_value

                for match in matches:
                    field_key = match.group(1)
                    placeholders.append(
                        ExcelPlaceholder(
                            sheet_name=sheet.title,
                            coordinate=cell.coordinate,
                            field_key=field_key,
                            raw_value=raw_value,
                        )
                    )

                    exact_placeholder = raw_value.strip() == match.group(0) and len(matches) == 1
                    if field_key not in values or values[field_key] is None:
                        if field_key not in missing_fields:
                            missing_fields.append(field_key)
                        replacement: Any = None if exact_placeholder else ""
                    else:
                        replacement = values[field_key]
                        replaced_count += 1

                    if exact_placeholder:
                        new_value = replacement
                    else:
                        replacement_text = "" if replacement is None else str(replacement)
                        new_value = str(new_value).replace(match.group(0), replacement_text)

                cell.value = new_value

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return ExcelFillResult(
        output_path=str(output),
        replaced_count=replaced_count,
        missing_fields=missing_fields,
        placeholders=placeholders,
    )
